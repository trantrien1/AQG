"""Export helpers for local question banks."""
from __future__ import annotations

import json
import re
import time
from pathlib import Path
from typing import Any, Dict, List

from . import bank_store

EXPORT_DIR = bank_store.DATA_DIR / 'bank_exports'


def _safe_name(value: str) -> str:
    value = re.sub(r'[^A-Za-z0-9_-]+', '_', value or 'question_bank').strip('_')
    return value[:80] or 'question_bank'


def _timestamp() -> str:
    return time.strftime('%Y%m%d_%H%M%S', time.localtime())


def _answer_text(question: Dict[str, Any]) -> str:
    answer_key = question.get('answer_key')
    for option in question.get('options') or []:
        if option.get('key') == answer_key:
            return str(option.get('text') or '')
    return ''


def _option_text(question: Dict[str, Any], key: str) -> str:
    for option in question.get('options') or []:
        if option.get('key') == key:
            return str(option.get('text') or '')
    return ''


def _rows(questions: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    rows = []
    for idx, question in enumerate(questions, start=1):
        source = question.get('source') or {}
        rows.append({
            'No': idx,
            'Question ID': question.get('question_id') or '',
            'Topic': question.get('topic') or '',
            'Bloom': question.get('cognitive_level') or '',
            'Difficulty': question.get('difficulty_target') or '',
            'Stem': question.get('stem') or '',
            'A': _option_text(question, 'A'),
            'B': _option_text(question, 'B'),
            'C': _option_text(question, 'C'),
            'D': _option_text(question, 'D'),
            'Answer Key': question.get('answer_key') or '',
            'Answer Text': _answer_text(question),
            'Explanation': question.get('short_explanation') or question.get('explanation') or '',
            'Source Pages': ', '.join(str(p) for p in (source.get('pages') or [])) if isinstance(source, dict) else '',
            'Source Quote': source.get('quote') if isinstance(source, dict) else '',
        })
    return rows


def export_bank(bank_id: str, fmt: str) -> Path:
    payload = bank_store.bank_export_payload(bank_id)
    if payload is None:
        raise KeyError('bank not found')
    bank = payload['bank']
    questions = payload['questions']
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    base = f"{_safe_name(bank.get('name') or bank_id)}_{_timestamp()}"
    fmt = fmt.lower().strip()
    if fmt == 'json':
        path = EXPORT_DIR / f'{base}.json'
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
        return path
    if fmt == 'xlsx':
        return _export_xlsx(EXPORT_DIR / f'{base}.xlsx', bank, questions)
    if fmt == 'docx':
        return _export_docx(EXPORT_DIR / f'{base}.docx', bank, questions)
    raise ValueError('format must be json|xlsx|docx')


def _export_xlsx(path: Path, bank: Dict[str, Any], questions: List[Dict[str, Any]]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Questions'
    rows = _rows(questions)
    headers = list(rows[0].keys()) if rows else [
        'No', 'Question ID', 'Topic', 'Bloom', 'Difficulty', 'Stem',
        'A', 'B', 'C', 'D', 'Answer Key', 'Answer Text', 'Explanation',
        'Source Pages', 'Source Quote',
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='E5E7EB')
    for row in rows:
        ws.append([row.get(h, '') for h in headers])
    for idx, header in enumerate(headers, start=1):
        width = 18
        if header in {'Stem', 'Explanation', 'Source Quote'}:
            width = 48
        ws.column_dimensions[get_column_letter(idx)].width = width

    meta = wb.create_sheet('Bank')
    meta.append(['Name', bank.get('name')])
    meta.append(['Description', bank.get('description')])
    meta.append(['Question Count', len(questions)])
    wb.save(path)
    return path


def _export_docx(path: Path, bank: Dict[str, Any], questions: List[Dict[str, Any]]) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading(str(bank.get('name') or 'Question Bank'), level=1)
    if bank.get('description'):
        doc.add_paragraph(str(bank.get('description')))
    for idx, question in enumerate(questions, start=1):
        doc.add_heading(f'Câu {idx}', level=2)
        doc.add_paragraph(str(question.get('stem') or ''))
        for key in ('A', 'B', 'C', 'D'):
            doc.add_paragraph(f'{key}. {_option_text(question, key)}')
        doc.add_paragraph(f'Đáp án: {question.get("answer_key") or ""}. {_answer_text(question)}')
        explanation = question.get('short_explanation') or question.get('explanation') or ''
        if explanation:
            doc.add_paragraph(f'Giải thích: {explanation}')
        source = question.get('source') or {}
        quote = source.get('quote') if isinstance(source, dict) else ''
        if quote:
            doc.add_paragraph(f'Nguồn: {quote}')
    doc.save(path)
    return path
